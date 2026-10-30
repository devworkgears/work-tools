"""CSV／Excel（.xlsx）の取り込みと、元の列とモックの項目の対応づけ。

元のファイルには列がたくさんある。モックが使う項目ごとに「どの見出しの列か」を ColumnMap で決め、
同じ見出しの並びのファイルには保存した ColumnMap をそのまま当てる（保存は store）。
"""

import codecs
import csv
import io
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from .model import BomLine, CostHead, CostRow, ImportWarning, Item, Version, Volume, mbom_to_ebom
from .normalize import to_code

VENDOR_DIR = Path(__file__).resolve().parent.parent / "vendor"


@dataclass(frozen=True)
class FieldSpec:
    name: str
    label: str
    required: bool


# どの表に、どの項目があるか（name はモックの中の名前・label は画面とエラーに出す名前）
TABLE_FIELDS: dict[str, tuple[FieldSpec, ...]] = {
    "item": (
        FieldSpec("code", "品番", True),
        FieldSpec("name", "名称", True),
        FieldSpec("specification", "記号・寸法", False),
        FieldSpec("kind", "区分", False),
        FieldSpec("lifecycle", "状態", False),
        FieldSpec("has_children", "子部品の有無", False),
    ),
    "bom_line": (
        FieldSpec("parent", "親品番", True),
        FieldSpec("level", "レベル", False),
        FieldSpec("child", "子品番", True),
        FieldSpec("quantity", "数量", True),
        FieldSpec("balloon", "風船番号", False),
    ),
    "cost": (
        FieldSpec("mbom_code", "MBOM 品番", True),
        FieldSpec("material", CostHead.MATERIAL.value, False),
        FieldSpec("direct", CostHead.DIRECT.value, False),
        FieldSpec("indirect", CostHead.INDIRECT.value, False),
        FieldSpec("adjustment", CostHead.ADJUSTMENT.value, False),
        FieldSpec("amortization", CostHead.AMORTIZATION.value, False),
        FieldSpec("manufacturing", CostHead.MANUFACTURING.value, False),
    ),
    "volume": (
        FieldSpec("product", "製品の品番", True),
        FieldSpec("annual", "年間台数", True),
    ),
}

COST_FIELDS: dict[str, CostHead] = {
    "material": CostHead.MATERIAL,
    "direct": CostHead.DIRECT,
    "indirect": CostHead.INDIRECT,
    "adjustment": CostHead.ADJUSTMENT,
    "amortization": CostHead.AMORTIZATION,
    "manufacturing": CostHead.MANUFACTURING,
}


class MissingColumnsError(ValueError):
    """必要な項目に当たる列が無い。missing は足りない項目の名前（label）の一覧。"""

    def __init__(self, table: str, missing: Sequence[str]) -> None:
        self.table = table
        self.missing = tuple(missing)
        super().__init__(f"{table} の必要な列が足りない: {'・'.join(self.missing)}")


def normalize_headers(headers: Iterable[Any]) -> tuple[str, ...]:
    """見出しを文字列にそろえ、前後の空白と末尾の空の見出しを落とす。"""
    result = ["" if h is None else str(h).strip() for h in headers]
    while result and result[-1] == "":
        result.pop()
    return tuple(result)


@dataclass(frozen=True)
class Table:
    """読み込んだ表。headers は見出しの並び、rows は見出しの次の行から（値は読んだままの型）。"""

    name: str
    headers: tuple[str, ...]
    rows: tuple[tuple[Any, ...], ...]


@dataclass(frozen=True)
class ColumnMap:
    """ファイルの形（見出しの並び）ごとの列の対応。columns は項目の name → 元の見出し。"""

    table: str
    headers: tuple[str, ...]
    columns: Mapping[str, str]

    def __post_init__(self) -> None:
        if self.table not in TABLE_FIELDS:
            raise ValueError(f"知らない表: {self.table!r}（{'・'.join(TABLE_FIELDS)} のどれか）")
        object.__setattr__(self, "headers", normalize_headers(self.headers))
        known = {f.name for f in TABLE_FIELDS[self.table]}
        unknown = set(self.columns) - known
        if unknown:
            raise ValueError(f"{self.table} に無い項目: {sorted(unknown)}")
        object.__setattr__(self, "columns", {k: str(v).strip() for k, v in self.columns.items()})
        self.check()

    def check(self, headers: Iterable[Any] | None = None) -> None:
        """必要な項目に、対応が無いか対応した見出しが無ければ MissingColumnsError。

        headers を省くと、この対応づけを作ったファイルの見出しで確かめる。
        """
        present = self.headers if headers is None else normalize_headers(headers)
        missing = tuple(
            f.label
            for f in TABLE_FIELDS[self.table]
            if f.required and self.columns.get(f.name) not in present
        )
        if missing:
            raise MissingColumnsError(self.table, missing)

    def matches(self, headers: Iterable[Any]) -> bool:
        return normalize_headers(headers) == self.headers

    def records(self, table: Table) -> Iterable[tuple[int, dict[str, Any]]]:
        """(ファイルの行番号, 項目の name → 値)。対応の無い項目・元のファイルの余分な列は持たない。"""
        index = {}
        for name, header in self.columns.items():
            if header in table.headers:
                index[name] = table.headers.index(header)
        for number, row in enumerate(table.rows, start=2):
            if all(v is None or (isinstance(v, str) and v.strip() == "") for v in row):
                continue
            yield number, {name: (row[i] if i < len(row) else None) for name, i in index.items()}


def decode_csv_bytes(data: bytes) -> str:
    """UTF-8（BOM あり・なし）と Shift_JIS（cp932）を見分けて読む。

    BOM があれば UTF-8、無ければ UTF-8 として読めるか試し、読めなければ cp932。
    """
    if data.startswith(codecs.BOM_UTF8):
        return data[len(codecs.BOM_UTF8):].decode("utf-8")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("cp932")


def read_csv(path: str | Path) -> Table:
    path = Path(path)
    return _csv_table(path.name, path.read_bytes())


def _csv_table(name: str, data: bytes) -> Table:
    text = decode_csv_bytes(data)
    rows = list(csv.reader(io.StringIO(text, newline="")))
    if not rows:
        return Table(name, (), ())
    return Table(name, normalize_headers(rows[0]), tuple(tuple(r) for r in rows[1:]))


def load_openpyxl() -> Any:
    """フォルダ内の vendor/ にあれば優先して読み、無ければ手元の Python の openpyxl を使う。"""
    if VENDOR_DIR.is_dir() and str(VENDOR_DIR) not in sys.path:
        sys.path.insert(0, str(VENDOR_DIR))
    import openpyxl

    return openpyxl


def read_xlsx(path: str | Path, sheet: str | None = None) -> Table:
    """1行目を見出しとして読む。sheet を省くと最初のシート。値は計算済みの値を読む。"""
    path = Path(path)
    openpyxl = load_openpyxl()
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = workbook[sheet] if sheet is not None else workbook.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return Table(path.name, (), ())
    return Table(path.name, normalize_headers(rows[0]), tuple(tuple(r) for r in rows[1:]))


def read_table(path: str | Path, sheet: str | None = None) -> Table:
    """拡張子で CSV と .xlsx を読み分ける。"""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return read_csv(path)
    if suffix in (".xlsx", ".xlsm"):
        return read_xlsx(path, sheet)
    raise ValueError(f"読めないファイルの種類: {path.name}（.csv・.xlsx）")


def xlsx_sheet_names(data: bytes) -> tuple[str, ...]:
    """.xlsx の中身（バイト列）のシート名の並び。"""
    workbook = load_openpyxl().load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return tuple(workbook.sheetnames)
    finally:
        workbook.close()


def read_table_bytes(name: str, data: bytes, sheet: str | None = None) -> Table:
    """画面から送られたファイルの中身を、ファイル名の拡張子で CSV と .xlsx に読み分ける。"""
    suffix = Path(name).suffix.lower()
    if suffix == ".csv":
        return _csv_table(name, data)
    if suffix in (".xlsx", ".xlsm"):
        workbook = load_openpyxl().load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            if sheet is not None and sheet not in workbook.sheetnames:
                raise ValueError(f"シート {sheet} は {name} に無い")
            ws = workbook[sheet] if sheet is not None else workbook.worksheets[0]
            rows = list(ws.iter_rows(values_only=True))
        finally:
            workbook.close()
        if not rows:
            return Table(name, (), ())
        return Table(name, normalize_headers(rows[0]), tuple(tuple(r) for r in rows[1:]))
    raise ValueError(f"読めないファイルの種類: {name}（.csv・.xlsx）")


@dataclass
class VersionImport:
    """1つの版に、何本ものファイルを取り込む。重複の畳みと数量の食い違いは Version が持つ。"""

    version: Version
    warnings: list[ImportWarning] = field(default_factory=list)

    def add_file(self, path: str | Path, column_map: ColumnMap, sheet: str | None = None) -> None:
        self.add_table(read_table(path, sheet), column_map)

    def add_table(self, table: Table, column_map: ColumnMap) -> None:
        column_map.check(table.headers)
        names = [n for n in self.version.source_name.split(", ") if n]
        if table.name not in names:
            names.append(table.name)
            self.version.source_name = ", ".join(names)
        add_row = getattr(self, f"_add_{column_map.table}")
        self_rows = 0
        for number, record in column_map.records(table):
            # 頂点の品目を「親＝子」の1行で出す構成表がある。構成ではないので、行ごとに警告せず数だけ残す
            if column_map.table == "bom_line" and _is_self_row(record):
                self_rows += 1
                continue
            try:
                add_row(table.name, number, record)
            except (ValueError, TypeError) as exc:
                self._warn(table.name, number, f"行を読み飛ばした: {exc}")
        if self_rows:
            self._warn(table.name, None, f"親と子が同じ行（頂点の行）を {self_rows} 行外した")

    def _warn(self, file: str, row: int | None, message: str) -> None:
        self.warnings.append(ImportWarning(file, row, message))

    def _add_item(self, file: str, number: int, r: dict[str, Any]) -> None:
        item = Item(
            r["code"],
            r.get("name"),
            specification=r.get("specification"),
            kind=r.get("kind"),
            lifecycle=r.get("lifecycle"),
            has_children=_blank_to(r.get("has_children"), 0),
        )
        kept = self.version.items.get(item.code)
        if kept is not None and kept != item:
            self._warn(file, number, f"品目 {item.code} が別の内容で重複したので後の行を採った")
        self.version.add_item(item)

    def _add_bom_line(self, file: str, number: int, r: dict[str, Any]) -> None:
        self.version.add_bom_lines([BomLine(r["parent"], r.get("balloon"), r["child"], r["quantity"])])

    def _add_cost(self, file: str, number: int, r: dict[str, Any]) -> None:
        mbom = to_code(r["mbom_code"], name="MBOM の品番")
        try:
            code = mbom_to_ebom(mbom)
        except ValueError:
            code = mbom
            self._warn(file, number, f"MBOM の品番が11桁でないので変換せずに取り込んだ: {mbom}")
        values = {head: r.get(name) for name, head in COST_FIELDS.items()}
        row = CostRow(code, values, mbom_code=mbom)
        kept = self.version.cost_row(code)
        if kept is not None and kept != row:
            self._warn(file, number, f"原価 {code} が別の内容で重複したので後の行を採った")
        self.version.add_cost(row)

    def _add_volume(self, file: str, number: int, r: dict[str, Any]) -> None:
        volume = Volume(r["product"], r["annual"])
        kept = self.version.volume(volume.product)
        if kept is not None and kept != volume:
            self._warn(file, number, f"年間台数 {volume.product} が別の値で重複したので後の行を採った")
        self.version.add_volume(volume)


def _is_self_row(record: Mapping[str, Any]) -> bool:
    try:
        return to_code(record.get("parent")) == to_code(record.get("child"))
    except (ValueError, TypeError):
        return False


def _blank_to(value: Any, default: Any) -> Any:
    if value is None or (isinstance(value, str) and value.strip() == ""):
        return default
    return value
