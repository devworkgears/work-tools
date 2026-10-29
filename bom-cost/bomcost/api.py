"""画面の API（JSON）。HTTP のハンドラはここの Api.handle を呼ぶだけで、ソケットを開かずにテストできる。

計算は計算の層（rollup・scenario・doublecount・mold）に任せる。金額と数量は Decimal の文字列で返し、
画面の JS は桁区切りの表示だけをする。未登録の原価は null で返し、0 とは分ける。
"""

import io
import json
import re
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any, Callable
from urllib.parse import parse_qs

from .doublecount import Judgement, judge_double_counting
from .importer import load_openpyxl
from .model import BomLine, CostHead, Version
from .mold import usage_per_unit
from .normalize import to_code, to_decimal, to_optional_code
from .rollup import CostCalculator, Scope
from .scenario import Change, ChangeKind, ChangeRecord, Scenario
from .store import Store

# 画面に並べる5費目（償却費は計算から外し、詳細に登録値だけ出す）
HEADS: tuple[CostHead, ...] = (
    CostHead.MATERIAL,
    CostHead.DIRECT,
    CostHead.INDIRECT,
    CostHead.ADJUSTMENT,
    CostHead.MANUFACTURING,
)
MODES = {"diff": "差分", "bulk": "一括"}
JSON_TYPE = "application/json; charset=utf-8"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# テーブルモードの列（キー・見出し・型）。型は Excel 書き出しで数値にするかどうかに使う
TABLE_COLUMNS: tuple[tuple[str, str, str], ...] = (
    ("level", "レベル", "int"),
    ("parent", "親品番", "text"),
    ("balloon", "風船番号", "text"),
    ("code", "品番", "text"),
    ("name", "名称", "text"),
    ("quantity", "数量", "number"),
    ("usage", "1台あたりの使用数", "number"),
    *((head.name, head.value, "money") for head in HEADS),
    ("verdict", "判定", "text"),
)
MISSING_TEXT = "未登録"


class ApiError(Exception):
    def __init__(self, status: int, message: str) -> None:
        super().__init__(message)
        self.status = status
        self.message = message


@dataclass
class Response:
    status: int
    body: bytes
    content_type: str = JSON_TYPE
    headers: dict[str, str] = field(default_factory=dict)

    def json(self) -> Any:
        return json.loads(self.body.decode("utf-8"))


def _json_response(data: Any, status: int = 200) -> Response:
    return Response(status, json.dumps(data, ensure_ascii=False).encode("utf-8"))


def error_response(status: int, message: str) -> Response:
    return _json_response({"error": message}, status)


def version_label(version: Version) -> str:
    """画面と同じ「D-123 名称（第2版）」の形。名称が空なら「D-123（第1版）」。"""
    number = version.dev_number or version.version_id
    name = f" {version.dev_name}" if version.dev_name else ""
    return f"{number}{name}（第{version.revision}版）"


def _s(value: Decimal | None) -> str | None:
    """指数表記にならない Decimal の文字列。未登録は None。"""
    return None if value is None else format(value, "f")


def _mul(a: Decimal, b: Decimal | None) -> Decimal | None:
    return None if b is None else a * b


def _head(value: Any, default: CostHead) -> CostHead:
    if value is None or value == "":
        return default
    try:
        head = CostHead[str(value)]
    except KeyError:
        raise ApiError(400, f"費目が分からない: {value}") from None
    if head is CostHead.AMORTIZATION:
        raise ApiError(400, "償却費は計算から外している（型費は別枠）")
    return head


def _mode(value: Any) -> str:
    mode = "diff" if value is None or value == "" else str(value)
    if mode not in MODES:
        raise ApiError(400, f"モードが分からない: {value}（diff か bulk）")
    return mode


def _flag(value: Any) -> bool:
    return str(value).lower() in ("1", "true", "yes") if value is not None else False


def _number(value: Any, name: str) -> Decimal:
    if value is None or value == "":
        raise ApiError(400, f"{name}が空")
    if not isinstance(value, (str, int, float)) or isinstance(value, bool):
        raise ApiError(400, f"{name}を数として読めない: {value!r}")
    return to_decimal(value, name=name)


def _text_field(body: dict[str, Any], key: str, name: str) -> str:
    value = body.get(key)
    if not isinstance(value, (str, int)) or isinstance(value, bool) or str(value).strip() == "":
        raise ApiError(400, f"{name}が空")
    return str(value).strip()


def _known_codes(version: Version) -> set[str]:
    codes = set(version.items) | {row.code for row in version.cost_rows}
    for line in version.lines:
        codes.add(line.parent)
        codes.add(line.child)
    return codes


def _require_known(version: Version, code: str, name: str) -> None:
    if code not in _known_codes(version):
        raise ApiError(400, f"{name} {code} は版 {version.version_id} に無い")


def _change_text(change: Change) -> str:
    qty = _s(change.quantity)
    if change.kind is ChangeKind.ADD:
        return f"追加 {change.part} ×{qty}"
    if change.kind is ChangeKind.REMOVE:
        return f"削除 {change.part} ×{qty}"
    if change.kind is ChangeKind.REPLACE and change.parent is not None:
        return f"{_place(change)} → {change.replacement} ×{qty}"
    if change.kind is ChangeKind.REPLACE:
        return f"{change.part} → {change.replacement}（×{qty}）"
    return f"員数変更 {_place(change)} → {qty}"


def _place(change: Change) -> str:
    """場所つきの置換の元の部品（風船番号があれば #番号 を前に）。"""
    return change.part if change.balloon is None else f"#{change.balloon} {change.part}"


class _View:
    """1つの試算を、モード・費目・風船番号の分け方を決めて見る。原価は費目ごとに1回だけ計算して使い回す。"""

    def __init__(
        self, scenario: Scenario, mode: str, head: CostHead, by_balloon: bool, amortization: bool = False
    ) -> None:
        self.scenario = scenario
        self.amortization = amortization
        self.version = scenario.version
        self.mode = mode
        self.head = head
        self.by_balloon = by_balloon
        self.diff = mode == "diff"
        self.root = scenario.product_code if self.diff else scenario.base
        self.overrides: dict[str, tuple[BomLine, ...]] = scenario.overrides() if self.diff else {}
        self._calcs: dict[CostHead, CostCalculator] = {}
        self._incomplete: dict[tuple[str, CostHead], bool] = {}
        self._registered: dict[tuple[str, CostHead], bool] = {}
        self._usage: dict[str, Decimal] | None = None
        self._judgements: dict[str, tuple[Judgement, ...]] = {}
        self.result = self.calc(head).calculate(self.root)
        # 差分の試算の製品は仮の品目なので、自分の登録値が無いのは当たり前（未登録に数えない）
        self.missing = tuple(c for c in self.result.missing if not (self.diff and c == self.root))
        self.total = self.cost(self.root)

    def calc(self, head: CostHead) -> CostCalculator:
        if head not in self._calcs:
            self._calcs[head] = CostCalculator(self.version, head, self.scenario.policy, self.overrides,
                                               allow_amortization=self.amortization)
        return self._calcs[head]

    def lines(self, code: str) -> tuple[BomLine, ...]:
        if code in self.overrides:
            return self.overrides[code]
        return self.version.children(code, self.by_balloon)

    def own_value(self, code: str, head: CostHead) -> Decimal | None:
        if self.diff and code == self.root:
            return None
        return self.version.cost_value(code, head)

    def cost(self, code: str, head: CostHead | None = None) -> Decimal | None:
        """計算の層の原価。足し上げた範囲に登録値が1つも無ければ、0 でなく未登録（None）にする。"""
        head = head or self.head
        value = self.calc(head).item_cost(code)
        return value if self._has_registered(code, head) else None

    def _has_registered(self, code: str, head: CostHead) -> bool:
        key = (code, head)
        if key not in self._registered:
            calc = self.calc(head)
            own = self.own_value(code, head)
            self._registered[key] = own is not None or any(
                self._has_registered(line.child, head) for line in calc.children(code)
            )
        return self._registered[key]

    def incomplete(self, code: str, head: CostHead | None = None) -> bool:
        """計算に使った範囲のどこかに未登録があるか（計算の層の足し上げと同じ道をたどる）。"""
        head = head or self.head
        key = (code, head)
        if key not in self._incomplete:
            calc = self.calc(head)
            own = self.own_value(code, head)
            lines = calc.children(code)
            if not lines:
                value = own is None
            elif own is not None and calc.policy.scope_for(code, head) is Scope.INCLUDES_CHILDREN:
                value = False
            else:
                is_virtual = self.diff and code == self.root
                value = (own is None and not is_virtual) or any(
                    self.incomplete(line.child, head) for line in lines
                )
            self._incomplete[key] = value
        return self._incomplete[key]

    def usage(self, code: str) -> Decimal | None:
        if self._usage is None:
            self._usage = usage_per_unit(self.version, self.root, self.overrides)
        if code == self.root:
            return Decimal(1)
        return self._usage.get(code)

    def judgements(self, code: str) -> tuple[Judgement, ...]:
        # 差分の試算の製品は仮の構成なので、版の登録値と比べる意味が無い
        if self.diff and code == self.root:
            return ()
        if code not in self._judgements:
            self._judgements[code] = judge_double_counting(self.version, [code])
        return self._judgements[code]

    def judgement(self, code: str, head: CostHead | None = None) -> Judgement | None:
        head = head or self.head
        return next((j for j in self.judgements(code) if j.head is head), None)

    def verdict(self, code: str, head: CostHead | None = None) -> str | None:
        j = self.judgement(code, head)
        return None if j is None else j.verdict.value

    def verdict_text(self, code: str) -> str:
        return "・".join(f"{j.head.value}: {j.verdict.value}" for j in self.judgements(code))

    def role(self, parent: str | None, line: BomLine | None) -> str:
        if line is None:
            return "試算の製品" if self.diff else "土台の製品"
        if not self.diff or parent != self.root:
            return ""
        if line.child == self.scenario.base and line.quantity > 0:
            return "土台"
        return "削除" if line.quantity < 0 else "追加"

    def name(self, code: str) -> str:
        item = self.version.items.get(code)
        if item is not None:
            return item.name
        return self.scenario.name if self.diff and code == self.root else ""

    def node(self, parent: str | None, line: BomLine | None, code: str | None = None) -> dict[str, Any]:
        """line が構成行なら子の行。line が None なら code の品目そのもの（省けば根、根の数量は 1）。"""
        if line is not None:
            code = line.child
            quantity: Decimal | None = line.quantity
        else:
            code = code or self.root
            quantity = Decimal(1) if code == self.root else None
        cost = self.cost(code)
        judgement = self.judgement(code)
        item = self.version.items.get(code)
        return {
            "code": code,
            "name": self.name(code),
            "kind": "" if item is None else item.kind,
            "parent": parent,
            "balloon": None if line is None else line.balloon,
            "quantity": _s(quantity),
            "cost": _s(cost),
            "amount": None if quantity is None else _s(_mul(quantity, cost)),
            "own_missing": self.own_value(code, self.head) is None,
            "incomplete": self.incomplete(code),
            "has_children": bool(self.lines(code)),
            "role": self.role(parent, line) if line is not None or code == self.root else "",
            "verdict": None if judgement is None else judgement.verdict.value,
            # 判定が比べた2つの値（登録値と、子の足し上げ・製造原価なら4費目の合計）
            "verdict_values": None if judgement is None else [_s(judgement.item_value), _s(judgement.compared_value)],
            "usage": _s(self.usage(code)),
        }

    def table_rows(self) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []

        def add(level: int, parent: str | None, line: BomLine | None, usage: Decimal) -> None:
            code = self.root if line is None else line.child
            row: dict[str, Any] = {
                "id": len(rows),
                "level": level,
                "parent": parent or "",
                "balloon": "" if line is None or line.balloon is None else line.balloon,
                "code": code,
                "name": self.name(code),
                "quantity": _s(Decimal(1) if line is None else line.quantity),
                "usage": _s(usage),
                "role": self.role(parent, line),
                "verdict": self.verdict_text(code),
            }
            for head in HEADS:
                row[head.name] = _s(self.cost(code, head))
            rows.append(row)

        def walk(code: str, level: int, multiplier: Decimal, path: frozenset[str]) -> None:
            for line in self.lines(code):
                if line.child in path:
                    raise ValueError(f"構成が循環している: {line.child}")
                usage = multiplier * line.quantity
                add(level, code, line, usage)
                walk(line.child, level + 1, usage, path | {line.child})

        add(0, None, None, Decimal(1))
        walk(self.root, 1, Decimal(1), frozenset({self.root}))
        return rows


Handler = Callable[..., Response]


class Api:
    """画面の API。store は SQLite の保存（Store）。1つのスレッドから呼ぶこと（sqlite3 の接続の制約）。"""

    def __init__(self, store: Store) -> None:
        self.store = store
        self._routes: list[tuple[str, re.Pattern[str], Handler]] = [
            ("GET", re.compile(r"/api/meta"), self.meta),
            ("GET", re.compile(r"/api/versions"), self.versions),
            ("GET", re.compile(r"/api/items"), self.items),
            ("GET", re.compile(r"/api/scenarios"), self.scenarios),
            ("POST", re.compile(r"/api/scenarios"), self.create_scenario),
            ("GET", re.compile(r"/api/scenarios/(\d+)"), self.scenario_summary),
            ("PUT", re.compile(r"/api/scenarios/(\d+)"), self.update_scenario),
            ("GET", re.compile(r"/api/scenarios/(\d+)/tree"), self.tree),
            ("GET", re.compile(r"/api/scenarios/(\d+)/item"), self.item),
            ("POST", re.compile(r"/api/scenarios/(\d+)/changes"), self.add_change),
            ("POST", re.compile(r"/api/scenarios/(\d+)/undo"), self.undo),
            ("GET", re.compile(r"/api/scenarios/(\d+)/table"), self.table),
            ("POST", re.compile(r"/api/scenarios/(\d+)/export"), self.export),
            ("GET", re.compile(r"/api/scenarios/(\d+)/compare"), self.compare),
        ]
        # 2段目の画面（取り込み・手入力・部品から見る・モデルチェンジ）。screens が api を読むので、ここで読む
        from .group import group_routes
        from .lookup import lookup_routes
        from .masterbook import default_opener
        from .screens import screen_routes

        # マスタのブックの置き場（None なら DB と同じフォルダ）と、書いたブックを開く関数（None なら開かない）
        self.master_dir: Path | None = None
        self.open_file: Callable[[Path], None] | None = default_opener()
        self._routes += [(m, re.compile(p), h)
                         for m, p, h in screen_routes(self) + group_routes(self) + lookup_routes(self)]

    # 入口 ------------------------------------------------------------------

    def handle(self, method: str, path: str, query: str = "", body: bytes = b"") -> Response:
        """1つの要求を処理する。入力の誤りは 400 番台と日本語のエラー文で返し、例外は外に出さない。"""
        try:
            matched_path = False
            for route_method, pattern, handler in self._routes:
                m = pattern.fullmatch(path)
                if m is None:
                    continue
                matched_path = True
                if route_method != method:
                    continue
                params = {k: v[0] for k, v in parse_qs(query, keep_blank_values=True).items()}
                args = [int(g) for g in m.groups()]
                if method in ("POST", "PUT"):
                    return handler(*args, params, _parse_body(body))
                return handler(*args, params)
            if matched_path:
                return error_response(405, f"{method} は使えない: {path}")
            return error_response(404, f"API が無い: {path}")
        except ApiError as exc:
            return error_response(exc.status, exc.message)
        except KeyError as exc:
            return error_response(404, str(exc.args[0]) if exc.args else "見つからない")
        except (ValueError, TypeError) as exc:
            return error_response(400, str(exc))
        except Exception as exc:  # 画面を止めないため、想定外の失敗も JSON で返す
            return error_response(500, f"サーバーの中で失敗した: {type(exc).__name__}: {exc}")

    # 一覧 ------------------------------------------------------------------

    def meta(self, params: dict[str, str]) -> Response:
        return _json_response(
            {
                "heads": [{"id": h.name, "label": h.value} for h in HEADS],
                "modes": [{"id": k, "label": v} for k, v in MODES.items()],
                "table_columns": [{"key": k, "label": label, "type": t} for k, label, t in TABLE_COLUMNS],
            }
        )

    def versions(self, params: dict[str, str]) -> Response:
        return _json_response(
            [
                {
                    "version_id": v.version_id,
                    "dev_number": v.dev_number,
                    "dev_name": v.dev_name,
                    "revision": v.revision,
                    "imported_at": None if v.imported_at is None else v.imported_at.isoformat(),
                    "source_name": v.source_name,
                    "note": v.note,
                    "warning_count": v.warning_count,
                }
                for v in self.store.list_versions()
            ]
        )

    def items(self, params: dict[str, str]) -> Response:
        """版（manual を重ねる）にある品番と名称。足す・替える品番の入力の候補に使う。"""
        version_id = params.get("version")
        if not version_id:
            raise ApiError(400, "版（version）が空")
        version = self._load_version(version_id)
        return _json_response(
            [
                {"code": code, "name": version.items[code].name if code in version.items else ""}
                for code in sorted(_known_codes(version))
            ]
        )

    def scenarios(self, params: dict[str, str]) -> Response:
        return _json_response(
            [
                {"scenario_id": s.scenario_id, "name": s.name, "base": s.base, "version_id": s.version_id}
                for s in self.store.list_scenarios()
            ]
        )

    # 試算 ------------------------------------------------------------------

    def create_scenario(self, params: dict[str, str], body: dict[str, Any]) -> Response:
        name = _text_field(body, "name", "試算の名前")
        base = to_code(_text_field(body, "base", "土台の製品の品番"))
        product = to_code(_text_field(body, "product_code", "試算の製品の品番"))
        version_id = _text_field(body, "version_id", "版")
        head = _head(body.get("head"), CostHead.MATERIAL)
        if base == product:
            raise ApiError(400, "試算の製品の品番は土台の製品と別の品番にする")
        version = self._load_version(version_id)
        _require_known(version, base, "土台の製品")
        scenario = Scenario(name, base, product, version, head, self.store.load_scope_policy())
        scenario.calculate()  # 構成の循環などをここで弾く
        self.store.save_scenario(scenario)
        return _json_response(self._summary(scenario), 201)

    def scenario_summary(self, scenario_id: int, params: dict[str, str]) -> Response:
        return _json_response(self._summary(self._load_scenario(scenario_id)))

    def update_scenario(self, scenario_id: int, params: dict[str, str], body: dict[str, Any]) -> Response:
        """計算に使う版・費目・名前を変える。変更の履歴は残し、今の原価を計算し直す。"""
        scenario = self._load_scenario(scenario_id)
        if body.get("version_id") not in (None, ""):
            scenario.version = self._load_version(str(body["version_id"]))
        if body.get("head") not in (None, ""):
            scenario.head = _head(body["head"], scenario.head)
        if body.get("name") not in (None, ""):
            scenario.name = _text_field(body, "name", "試算の名前")
        scenario.calculate()
        self.store.save_scenario(scenario)
        return _json_response(self._summary(scenario))

    def add_change(self, scenario_id: int, params: dict[str, str], body: dict[str, Any]) -> Response:
        scenario = self._load_scenario(scenario_id)
        self.apply_change(scenario, body)
        self.store.save_scenario(scenario)
        return _json_response(self._summary(scenario))

    def apply_change(self, scenario: Scenario, body: dict[str, Any]) -> None:
        """body の変更1回分を確かめて試算に当てる（保存はしない）。通らなければ ApiError で、試算は変えない。"""
        kind_name = body.get("kind")
        try:
            kind = ChangeKind[str(kind_name)]
        except KeyError:
            raise ApiError(400, f"変更の種類が分からない: {kind_name}（ADD・REMOVE・REPLACE・SET_QUANTITY）") from None
        part = to_code(_text_field(body, "part", "品番"))
        if part == scenario.product_code:
            raise ApiError(400, "試算の製品そのものは追加・削除・置換の対象にできない")
        _require_known(scenario.version, part, "品番")
        parent = to_optional_code(body.get("parent"), name="親品番")
        balloon = to_optional_code(body.get("balloon"), name="風船番号")
        if body.get("replacements") is not None and kind is not ChangeKind.REPLACE:
            raise ApiError(400, "代わりの部品の並び（replacements）は置換だけに付ける")
        if parent is not None and kind not in (ChangeKind.REPLACE, ChangeKind.SET_QUANTITY):
            raise ApiError(400, "場所（parent・balloon）は置換だけに付ける（員数変更も可）")
        if balloon is not None and parent is None:
            raise ApiError(400, "風船番号（balloon）には親品番（parent）が要る")
        if kind is ChangeKind.REPLACE:
            pairs = self._replacements(scenario, body)
        else:
            pairs = [(None, _number(body.get("quantity", "1"), "数量"))]
        if parent is not None:
            location = (parent, balloon, part)
            label = f"{parent} の {part}" if balloon is None else f"{parent} の #{balloon} {part}"
            if parent != scenario.base and not usage_per_unit(scenario.version, scenario.base).get(parent):
                raise ApiError(400, f"親 {parent} は土台 {scenario.base} の構成に無い")
            if scenario.location_quantity(location) == 0:
                raise ApiError(400, f"{label} は版 {scenario.version.version_id} の構成に無い")
            if location in scenario.replaced_locations():
                raise ApiError(400, f"{label} はもう置換した（員数変更も含む）")
            # 場所の変更は元の版の使用数で倍率を取るので、親子にまたがる2つの変更は原価が誤る
            below = lambda code: {code, *usage_per_unit(scenario.version, code)}
            for done_parent, _, done_part in scenario.replaced_locations():
                if parent in below(done_part) or done_parent in below(part):
                    raise ApiError(400, f"{label} は変更済みの {done_parent} の {done_part} と親子に重なる（先に取り消す）")
        elif len(pairs) > 1:
            raise ApiError(400, "代わりの部品を複数入れるには親品番（parent）が要る")
        changes = [Change(kind, part, qty, code, parent, balloon) for code, qty in pairs]
        volume = scenario.version.volume(scenario.product_code)
        scenario.apply_batch(changes, volume=None if volume is None else volume.annual)

    @staticmethod
    def _replacements(scenario: Scenario, body: dict[str, Any]) -> list[tuple[str, Decimal]]:
        """代わりの部品（品番・員数）の並び。replacements が無ければ replacement・quantity の1つ。"""
        raw = body.get("replacements")
        if raw is None:
            raw = [{"code": body.get("replacement"), "quantity": body.get("quantity", "1")}]
        if not isinstance(raw, list) or not raw or not all(isinstance(r, dict) for r in raw):
            raise ApiError(400, "代わりの部品の並び（replacements）は品番・員数の組の配列にする")
        pairs = []
        for r in raw:
            code = to_code(_text_field({"code": r.get("code", r.get("replacement"))}, "code", "代わりの品番"))
            _require_known(scenario.version, code, "代わりの品番")
            if code == scenario.product_code:
                raise ApiError(400, "試算の製品そのものを代わりの品番にはできない")
            pairs.append((code, _number(r.get("quantity", "1"), "数量")))
        return pairs

    def undo(self, scenario_id: int, params: dict[str, str], body: dict[str, Any]) -> Response:
        scenario = self._load_scenario(scenario_id)
        if scenario.undo() is None:
            raise ApiError(400, "取り消す変更が無い")
        self.store.save_scenario(scenario)
        return _json_response(self._summary(scenario))

    # 見る ------------------------------------------------------------------

    def tree(self, scenario_id: int, params: dict[str, str]) -> Response:
        """code（省けば根）の行と、その直下の子の行。ツリーは開いた所だけ取りに来る。"""
        return _json_response(self.tree_payload(self._view(scenario_id, params), params.get("code")))

    @classmethod
    def tree_payload(cls, view: _View, code: str | None) -> dict[str, Any]:
        """見方 view での code（省けば根）の行と、その直下の子の行。"""
        code = to_code(code or view.root)
        if code != view.root and code not in _known_codes(view.version):
            raise ApiError(400, f"品番 {code} は版 {view.version.version_id} に無い")
        node = view.node(None, None, code)
        children = [view.node(code, line) for line in view.lines(code)]
        return {"view": cls._view_info(view), "node": node, "children": children}

    def item(self, scenario_id: int, params: dict[str, str]) -> Response:
        """選んだ品目の詳細: 費目5つ（登録値と計算値）・出どころ・二重計上の判定。"""
        view = self._view(scenario_id, params)
        code = to_code(params.get("code") or view.root)
        version = view.version
        if code != view.root and code not in _known_codes(version):
            raise ApiError(400, f"品番 {code} は版 {version.version_id} に無い")
        item = version.items.get(code)
        row = version.cost_row(code)
        heads = []
        for head in HEADS:
            judgement = next((j for j in view.judgements(code) if j.head is head), None)
            heads.append(
                {
                    "id": head.name,
                    "label": head.value,
                    "registered": _s(view.own_value(code, head)),
                    "calculated": _s(view.cost(code, head)),
                    "incomplete": view.incomplete(code, head),
                    "verdict": None if judgement is None else judgement.verdict.value,
                    "compared": None if judgement is None else _s(judgement.compared_value),
                }
            )
        root_codes = {line.child for line in view.lines(view.root)}
        return _json_response(
            {
                "view": self._view_info(view),
                "code": code,
                "name": view.name(code),
                "kind": "" if item is None else item.kind,
                "lifecycle": "" if item is None else item.lifecycle,
                "item_source": None if item is None else item.source.value,
                "cost_source": None if row is None else row.source.value,
                "has_children": bool(view.lines(code)),
                "is_root": code == view.root,
                "in_root": code in root_codes,
                "usage": _s(view.usage(code)),
                "amortization": _s(version.cost_value(code, CostHead.AMORTIZATION)),
                "heads": heads,
            }
        )

    def table(self, scenario_id: int, params: dict[str, str]) -> Response:
        view = self._view(scenario_id, params)
        return _json_response(
            {
                "view": self._view_info(view),
                "columns": [{"key": k, "label": label, "type": t} for k, label, t in TABLE_COLUMNS],
                "rows": view.table_rows(),
            }
        )

    def export(self, scenario_id: int, params: dict[str, str], body: dict[str, Any]) -> Response:
        """テーブルモードの表示中の行・列を .xlsx で返す。行は row_ids（表の id）の並び、省けば全部。"""
        view = self._view(scenario_id, {k: body.get(k) for k in ("mode", "head", "by_balloon")})
        rows = view.table_rows()
        columns_by_key = {k: (label, t) for k, label, t in TABLE_COLUMNS}
        keys = body.get("columns") or [k for k, _, _ in TABLE_COLUMNS]
        if not isinstance(keys, list) or any(k not in columns_by_key for k in keys):
            raise ApiError(400, f"列が分からない: {keys}")
        ids = body.get("row_ids")
        if ids is not None:
            if not isinstance(ids, list) or any(
                not isinstance(i, int) or isinstance(i, bool) or not 0 <= i < len(rows) for i in ids
            ):
                raise ApiError(400, "行の指定（row_ids）が表の行と合わない")
            rows = [rows[i] for i in ids]
        data = _workbook(view, keys, columns_by_key, rows)
        return Response(
            200,
            data,
            XLSX_TYPE,
            {"Content-Disposition": f'attachment; filename="bom-cost-{scenario_id}.xlsx"'},
        )

    def compare(self, scenario_id: int, params: dict[str, str]) -> Response:
        """今の試算を別の版で計算し直し、直下の品目ごとの差を返す。"""
        scenario = self._load_scenario(scenario_id)
        other_id = params.get("version")
        if not other_id:
            raise ApiError(400, "比べる版（version）が空")
        other = self._load_version(other_id)
        head = _head(params.get("head"), scenario.head)
        cmp = scenario.compare_versions(scenario.version, other, head)
        view = _View(scenario, "diff", head, False)
        lines = []
        for row, line in zip(cmp.lines, scenario.root_lines()):
            lines.append(
                {
                    "part": row.part,
                    "name": view.name(row.part) or (other.items[row.part].name if row.part in other.items else ""),
                    "role": view.role(view.root, line),
                    "quantity": _s(row.quantity),
                    "amount_a": _s(row.amount_a),
                    "amount_b": _s(row.amount_b),
                    "difference": _s(row.difference),
                }
            )
        return _json_response(
            {
                "scenario_id": scenario_id,
                "product_code": scenario.product_code,
                "version_a": cmp.version_a,
                "version_b": cmp.version_b,
                "head": head.name,
                "head_label": head.value,
                "lines": lines,
                "total_a": _s(cmp.total_a),
                "total_b": _s(cmp.total_b),
                "difference": _s(cmp.difference),
            }
        )

    # 中身 ------------------------------------------------------------------

    def _load_scenario(self, scenario_id: int) -> Scenario:
        try:
            return self.store.load_scenario(scenario_id)
        except KeyError:
            raise ApiError(404, f"試算 {scenario_id} は無い") from None

    def _load_version(self, version_id: str) -> Version:
        try:
            return self.store.load_version(version_id)
        except KeyError:
            raise ApiError(404, f"版 {version_id} は無い") from None

    def _view(self, scenario_id: int, params: dict[str, Any]) -> _View:
        scenario = self._load_scenario(scenario_id)
        return _View(
            scenario,
            _mode(params.get("mode")),
            _head(params.get("head"), scenario.head),
            _flag(params.get("by_balloon")),
        )

    @staticmethod
    def _view_info(view: _View) -> dict[str, Any]:
        return {
            "mode": view.mode,
            "mode_label": MODES[view.mode],
            "root": view.root,
            "version_id": view.version.version_id,
            "head": view.head.name,
            "head_label": view.head.value,
            "by_balloon": view.by_balloon,
            "total": _s(view.total),
            "missing": list(view.missing),
            "premises": _premises(view.version.version_id, view.head, bool(view.missing)),
        }

    def _summary(self, scenario: Scenario) -> dict[str, Any]:
        view = _View(scenario, "diff", scenario.head, False)
        result = view.result
        start = scenario.calculate(changes=())
        volume = scenario.version.volume(scenario.product_code)
        history = [_record_json(r) for r in scenario.history]
        last = history[-1] if history else None
        return {
            "scenario": {
                "scenario_id": scenario.scenario_id,
                "name": scenario.name,
                "base": scenario.base,
                "product_code": scenario.product_code,
                "version_id": scenario.version.version_id,
                "head": scenario.head.name,
                "head_label": scenario.head.value,
                "note": scenario.note,
            },
            "total": _s(view.total),
            "start": {"base": scenario.base, "cost": _s(start.total)},
            "missing": list(view.missing),
            "conflicts": [
                {"parent": c.key[0], "balloon": c.key[1], "child": c.key[2], "quantities": [_s(q) for q in c.quantities]}
                for c in result.conflicts
            ],
            "volume": None if volume is None else {"annual": _s(volume.annual), "kind": volume.kind.value},
            "history": history,
            "last": None
            if last is None
            else {k: last[k] for k in ("cost_before", "cost_after", "difference", "text", "version_id", "head_label")},
            "premises": _premises(scenario.version.version_id, scenario.head, bool(view.missing)),
        }


def _premises(version_id: str, head: CostHead, missing: bool) -> dict[str, Any]:
    """結果と履歴に必ず添える前提。"""
    return {
        "version_id": version_id,
        "head": head.name,
        "head_label": head.value,
        "reference": "参考値",
        "amortization_excluded": True,
        "has_missing": missing,
    }


def _record_json(record: ChangeRecord) -> dict[str, Any]:
    change = record.change
    return {
        "sequence": record.sequence,
        "kind": change.kind.name,
        "kind_label": change.kind.value,
        "part": change.part,
        "replacement": change.replacement,
        "quantity": _s(change.quantity),
        "text": _change_text(change),
        "cost_before": _s(record.cost_before),
        "cost_after": _s(record.cost_after),
        "difference": _s(record.difference),
        "version_id": record.version_id,
        "head": record.head.name,
        "head_label": record.head.value,
        "volume": _s(record.volume),
        "allocation": None if record.allocation is None else record.allocation.value,
        "recorded_at": record.recorded_at.isoformat(timespec="seconds"),
        "parent": change.parent,
        "balloon": change.balloon,
        "batch": record.batch,
    }


def _parse_body(body: bytes) -> dict[str, Any]:
    if not body:
        return {}
    try:
        data = json.loads(body.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError):
        raise ApiError(400, "送られた内容を JSON として読めない") from None
    if not isinstance(data, dict):
        raise ApiError(400, "送る内容は JSON のオブジェクトにする")
    return data


def _workbook(
    view: _View, keys: list[str], columns: dict[str, tuple[str, str]], rows: list[dict[str, Any]]
) -> bytes:
    openpyxl = load_openpyxl()
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "構成"
    ws.append([columns[k][0] for k in keys])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        values: list[Any] = []
        for key in keys:
            value = row.get(key)
            kind = columns[key][1]
            if kind in ("number", "money"):
                # 未登録は空欄にすると 0 と区別がつかないので、文字で残す
                values.append(MISSING_TEXT if value is None else Decimal(value))
            else:
                values.append(value)
        ws.append(values)
    ws.freeze_panes = "A2"

    premises = wb.create_sheet("前提")
    scenario = view.scenario
    for label, value in (
        ("試算", scenario.name),
        ("モード", MODES[view.mode]),
        ("根の品番", view.root),
        ("計算に使った版", version_label(view.version)),
        ("費目（ツリー・合計）", view.head.value),
        ("値の扱い", "参考値"),
        ("償却費", "計算から外す（型費は別枠）"),
        ("未登録", "あり（" + "・".join(view.missing) + "）" if view.missing else "なし"),
        ("合計", MISSING_TEXT if view.total is None else view.total),
    ):
        premises.append([label, value])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
