"""マスタ登録（製品・部品・年間台数・型費・流用先・部品の対応表）を Excel のブックで書き出し、取り込む。

取り込みは追加と上書きだけ（ブックから行を消しても登録は消さない）。確認（plan_book）は何も保存せず、
決定（apply_plan）は全部を1回の書き込みで入れる。列の見出しは画面の表にそろえる。
"""

import io
import os
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any, Callable

from .api import HEADS
from .importer import load_openpyxl
from .model import MANUAL_VERSION_ID, CostHead, CostRow, Item, Version, Volume
from .mold import MoldCost, ReuseTarget
from .normalize import to_code, to_decimal, to_symbol
from .store import Store

PRODUCT = "製品"
PART = "部品"
VOLUME = "年間台数"
MOLD = "型費"
REUSE = "流用先"
PART_MAP = "部品の対応表"

HEAD_LABELS = tuple(h.value for h in HEADS)
SHEETS: dict[str, tuple[str, ...]] = {
    PRODUCT: ("品番", "名称", *HEAD_LABELS),
    PART: ("品番", "名称", "区分", *HEAD_LABELS),
    VOLUME: ("製品の品番", "年間台数"),
    MOLD: ("部品の品番", "型費（円）", "償却年数"),
    REUSE: ("部品", "流用先の製品", "使用数（個/台）", "置き換わる古い部品"),
    PART_MAP: ("元の製品", "新しい製品", "元の部品", "元の風船番号", "新しい部品", "新しい風船番号"),
}
DEFAULT_YEARS = Decimal(2)

PartMapKey = tuple[str, str, str, str, str | None, str | None]


@dataclass(frozen=True)
class BookError:
    sheet: str
    row: int | None  # Excel の行番号（見出しが1行目）。シート全体の誤りは None
    message: str


@dataclass
class SheetCount:
    sheet: str
    added: int = 0
    changed: int = 0
    unchanged: int = 0


@dataclass
class MasterPlan:
    """ブックと今の登録の差。errors が空のときだけ apply_plan で入れられる。"""

    counts: dict[str, SheetCount] = field(default_factory=lambda: {s: SheetCount(s) for s in SHEETS})
    errors: list[BookError] = field(default_factory=list)
    items: dict[str, Item] = field(default_factory=dict)
    costs: dict[str, CostRow | None] = field(default_factory=dict)  # None は原価の行を外す
    volumes: list[Volume] = field(default_factory=list)
    molds: list[MoldCost] = field(default_factory=list)
    reuse: list[ReuseTarget] = field(default_factory=list)
    part_maps: list[PartMapKey] = field(default_factory=list)
    dev_number: str = ""  # 入れる先の開発番号（空は共通）

    def total(self, key: str) -> int:
        return sum(getattr(c, key) for c in self.counts.values())


class _RowError(ValueError):
    pass


@dataclass
class _Current:
    """今の登録（行ごとに読み直さないよう1回だけ読む）。"""

    manual: Version
    molds: dict[str, MoldCost]
    reuse: list[ReuseTarget]
    part_maps: set[PartMapKey]

    @classmethod
    def load(cls, store: Store, dev_number: str = "") -> "_Current":
        return cls(
            store.load_manual(dev_number),
            store.mold_costs(dev_number),
            store.reuse_targets(dev_number=dev_number),
            {(old, new, r.old_part, r.new_part, r.old_balloon, r.new_balloon)
             for old, new in store.part_map_pairs(dev_number) for r in store.part_map_rows(old, new, dev_number)},
        )


# 書き出し ------------------------------------------------------------------


def book_bytes(store: Store, product_kind: str, dev_number: str = "") -> bytes:
    """今の登録（その開発番号の分・空は共通）を6シートのブックにする。

    試算（検討案でできた製品）は手入力の版に無いので入らない。
    """
    openpyxl = load_openpyxl()
    from openpyxl.styles import Font

    manual = store.load_manual(dev_number)
    codes = list(manual.items) + [r.code for r in manual.cost_rows if r.code not in manual.items]
    products: list[list[Any]] = []
    parts: list[list[Any]] = []
    for code in codes:
        item = manual.items.get(code)
        name = "" if item is None else item.name
        kind = "" if item is None else item.kind
        costs = [manual.cost_value(code, h) for h in HEADS]
        if kind == product_kind:
            products.append([code, name, *costs])
        else:
            parts.append([code, name, kind, *costs])
    rows: dict[str, list[list[Any]]] = {
        PRODUCT: products,
        PART: parts,
        VOLUME: [[v.product, v.annual] for v in manual.volumes.values()],
        MOLD: [[m.part, m.amount, m.years] for m in store.mold_costs(dev_number).values()],
        REUSE: [[t.part, t.product, t.quantity_per_unit, t.replaced_part]
                for t in store.reuse_targets(dev_number=dev_number)],
        PART_MAP: [[old, new, r.old_part, r.old_balloon, r.new_part, r.new_balloon]
                   for old, new in store.part_map_pairs(dev_number)
                   for r in store.part_map_rows(old, new, dev_number)],
    }
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet, headers in SHEETS.items():
        ws = wb.create_sheet(sheet)
        ws.append(list(headers))
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows[sheet]:
            ws.append(row)
        ws.freeze_panes = "A2"
        for i, header in enumerate(headers):
            ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = max(12, len(header) * 2 + 2)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def save_book(path: Path, data: bytes) -> None:
    """一時ファイルに書いてから置き換える。Excel で開いたままなら PermissionError が出る（Windows）。"""
    temp = path.with_name(path.name + ".tmp")
    temp.write_bytes(data)
    try:
        os.replace(temp, path)
    finally:
        if temp.exists():
            temp.unlink()


def default_opener() -> Callable[[Path], None] | None:
    """既定のアプリで開く関数。os.startfile は Windows にしか無いので、他では None（開かない）。"""
    return getattr(os, "startfile", None)


# 取り込み ------------------------------------------------------------------


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _code(value: Any, name: str) -> str:
    if value is None or _text(value) == "":
        raise _RowError(f"{name}が空")
    return to_code(value, name=name)


def _optional_code(value: Any, name: str) -> str | None:
    text = _text(value)
    return None if text == "" else to_code(text, name=name)


def _decimal(value: Any, name: str, required: bool) -> Decimal | None:
    if value is None or (isinstance(value, str) and value.strip() == ""):
        if required:
            raise _RowError(f"{name}が空")
        return None
    return to_decimal(value, name=name)


def _read_sheets(data: bytes) -> dict[str, list[tuple[Any, ...]]]:
    workbook = load_openpyxl().load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return {ws.title: list(ws.iter_rows(values_only=True)) for ws in workbook.worksheets}
    finally:
        workbook.close()


def plan_book(store: Store, product_kind: str, data: bytes, dev_number: str = "") -> MasterPlan:
    """ブックを読み、今の登録（その開発番号の分・空は共通）との差（追加・変更・同じ）と行の誤りを返す。何も保存しない。

    読めないブックは ValueError 以外の例外（zipfile など）も出るので、呼ぶ側でまとめて受ける。
    """
    sheets = _read_sheets(data)
    plan = MasterPlan(dev_number=dev_number)
    current = _Current.load(store, dev_number)
    item_rows: dict[str, tuple[str, int]] = {}  # 品番 → 最初に出たシートと行（製品と部品をまたいで見る）

    for sheet, headers in SHEETS.items():
        if sheet not in sheets:
            plan.errors.append(BookError(sheet, None, "シートが無い"))
            continue
        rows = sheets[sheet]
        head = [_text(v) for v in rows[0]] if rows else []
        missing = [h for h in headers if h not in head]
        if missing:
            plan.errors.append(BookError(sheet, 1, f"列が無い: {'・'.join(missing)}"))
            continue
        index = {h: head.index(h) for h in headers}
        seen: dict[Any, int] = {}
        for number, raw in enumerate(rows[1:], start=2):
            values = {h: (raw[index[h]] if index[h] < len(raw) else None) for h in headers}
            if all(_text(v) == "" for v in values.values()):
                continue
            try:
                key, status = _plan_row(plan, current, product_kind, sheet, values, item_rows, number)
                if key in seen:
                    raise _RowError(f"同じ行が {seen[key]} 行目にもある")
                seen[key] = number
            except (_RowError, ValueError, TypeError) as exc:
                plan.errors.append(BookError(sheet, number, str(exc)))
                continue
            count = plan.counts[sheet]
            setattr(count, status, getattr(count, status) + 1)
    return plan


def _plan_row(
    plan: MasterPlan,
    current: _Current,
    product_kind: str,
    sheet: str,
    v: dict[str, Any],
    item_rows: dict[str, tuple[str, int]],
    number: int,
) -> tuple[Any, str]:
    """1行を読んで plan に足し、(重なりを見るキー, added・changed・unchanged) を返す。"""
    manual = current.manual
    if sheet in (PRODUCT, PART):
        code = _code(v["品番"], "品番")
        name = _text(v["名称"])
        if not name:
            raise _RowError("名称が空")
        if sheet == PRODUCT:
            kind = product_kind
        else:
            kind = to_symbol(_text(v["区分"]) or None, name="区分")
            if kind == product_kind:
                raise _RowError(f"部品の区分に製品の記号 {kind} は使えない")
        heads = {h: _decimal(v[h.value], h.value, False) for h in HEADS}
        if code in item_rows and item_rows[code][0] != sheet:
            other, row = item_rows[code]
            raise _RowError(f"品番 {code} は{other}のシートの {row} 行目にもある")
        kept = manual.items.get(code)
        old_cost = manual.cost_row(code)
        values: dict[CostHead, Decimal | None] = {h: None for h in CostHead}
        if old_cost is not None:
            values.update(old_cost.values)
        values.update(heads)
        new_cost = (CostRow(code, values, mbom_code=None if old_cost is None else old_cost.mbom_code)
                    if any(x is not None for x in values.values()) else None)
        item = Item(code, name, kind, "" if kept is None else kept.lifecycle,
                    False if kept is None else kept.has_children)
        if kept is None and old_cost is None:
            status = "added"
        elif (kept is not None and kept.name == name and kept.kind == kind
              and all(manual.cost_value(code, h) == heads[h] for h in HEADS)):
            status = "unchanged"
        else:
            status = "changed"
        if code not in item_rows:
            item_rows[code] = (sheet, number)
            if status != "unchanged":
                plan.items[code] = item
                plan.costs[code] = new_cost
        return code, status

    if sheet == VOLUME:
        product = _code(v["製品の品番"], "製品の品番")
        annual = _decimal(v["年間台数"], "年間台数", True)
        assert annual is not None
        if annual < 0:
            raise _RowError("年間台数は 0 以上にする")
        old = manual.volumes.get(product)
        status = "added" if old is None else ("unchanged" if old.annual == annual else "changed")
        if status != "unchanged" and product not in {x.product for x in plan.volumes}:
            plan.volumes.append(Volume(product, annual))
        return product, status

    if sheet == MOLD:
        part = _code(v["部品の品番"], "部品の品番")
        amount = _decimal(v["型費（円）"], "型費", True)
        years = _decimal(v["償却年数"], "償却年数", False) or DEFAULT_YEARS
        assert amount is not None
        if amount < 0:
            raise _RowError("型費は 0 以上にする")
        if years <= 0:
            raise _RowError("償却年数は正の数にする")
        old_mold = current.molds.get(part)
        status = ("added" if old_mold is None
                  else "unchanged" if (old_mold.amount, old_mold.years) == (amount, years) else "changed")
        if status != "unchanged" and part not in {m.part for m in plan.molds}:
            plan.molds.append(MoldCost(part, amount, years))
        return part, status

    if sheet == REUSE:
        part = _code(v["部品"], "部品の品番")
        product = _code(v["流用先の製品"], "流用先の品番")
        quantity = _decimal(v["使用数（個/台）"], "使用数", True)
        assert quantity is not None
        if quantity <= 0:
            raise _RowError("使用数は正の数にする")
        replaced = _optional_code(v["置き換わる古い部品"], "置き換わる古い部品")
        old_targets = [t for t in current.reuse if t.part == part and t.product == product]
        if not old_targets:
            status = "added"
        elif len(old_targets) == 1 and (old_targets[0].quantity_per_unit, old_targets[0].replaced_part) == (
                quantity, replaced):
            status = "unchanged"
        else:
            status = "changed"
        key = (part, product)
        if status != "unchanged" and key not in {(t.part, t.product) for t in plan.reuse}:
            plan.reuse.append(ReuseTarget(part, product, quantity, replaced))
        return key, status

    map_key: PartMapKey = (
        _code(v["元の製品"], "元の製品の品番"),
        _code(v["新しい製品"], "新しい製品の品番"),
        _code(v["元の部品"], "元の部品の品番"),
        _code(v["新しい部品"], "新しい部品の品番"),
        _optional_code(v["元の風船番号"], "風船番号"),
        _optional_code(v["新しい風船番号"], "風船番号"),
    )
    if map_key[0] == map_key[1]:
        raise _RowError("元の製品と新しい製品は別の品番にする")
    if map_key[2] == map_key[3]:
        raise _RowError("同じ品番どうしは対応表に入れなくても自動で対応する")
    status = "unchanged" if map_key in current.part_maps else "added"
    if status == "added" and map_key not in plan.part_maps:
        plan.part_maps.append(map_key)
    return map_key, status


def apply_plan(store: Store, plan: MasterPlan) -> None:
    """誤りの無い plan を1回の書き込みで入れる（plan の開発番号の登録に入る）。"""
    if plan.errors:
        raise ValueError("行の誤りがある plan は入れない")
    manual = store.load_manual(plan.dev_number)
    new_manual: Version | None = None
    if plan.items or plan.volumes:
        new_manual = Version(MANUAL_VERSION_ID)
        items = dict(manual.items)
        items.update(plan.items)
        new_manual.add_items(items.values())
        new_manual.add_bom_lines(manual.lines)
        for conflict in manual.conflicts:
            new_manual.add_conflict(conflict.key, conflict.quantities)
        costs: dict[str, CostRow] = {r.code: r for r in manual.cost_rows}
        for code, row in plan.costs.items():
            if row is None:
                costs.pop(code, None)
            else:
                costs[code] = row
        new_manual.add_costs(costs.values())
        volumes = dict(manual.volumes)
        volumes.update({x.product: x for x in plan.volumes})
        for volume in volumes.values():
            new_manual.add_volume(volume)
    store.apply_master(new_manual, plan.molds, plan.reuse, plan.part_maps, plan.dev_number)


def plan_json(plan: MasterPlan) -> dict[str, Any]:
    return {
        "added": plan.total("added"),
        "changed": plan.total("changed"),
        "unchanged": plan.total("unchanged"),
        "sheets": [{"sheet": c.sheet, "added": c.added, "changed": c.changed, "unchanged": c.unchanged}
                   for c in plan.counts.values()],
        "errors": [{"sheet": e.sheet, "row": e.row, "message": e.message} for e in plan.errors],
    }
