"""③原価比較の表（compare/table の返事）を .xlsx に。並びは画面の表と同じで、前提のシートを付ける。

数は数値のセルで入れ、未登録は空のセルに印（塗りとコメント）を付ける。償却費を無視した列と、現行が 0 で出せない差% は「—」。
"""

import io
import re
from datetime import datetime
from decimal import Decimal
from typing import Any

from .importer import load_openpyxl

SHEET = "比較"
PREMISES = "前提"
DASH = "—"
MISSING_NOTE = "未登録・計から除く"
EXCLUDED_NOTE = "計から除く"
MONEY_FORMAT = "#,##0"
MONEY_FRACTION_FORMAT = "#,##0.00"
PERCENT_FORMAT = "0.0%"
MARK_COLOR = "FFF2CC"
_UNSAFE = re.compile(r'[\\/:*?"<>|\x00-\x1f]')


def compare_file_name(version_id: str, names: list[str]) -> str:
    """版と検討案が分かるファイル名。ファイル名に使えない文字は _ に。"""
    return _UNSAFE.sub("_", f"原価比較_{version_id}_{'・'.join(names)}") + ".xlsx"


def compare_workbook(table: dict[str, Any], version_text: str, query: str, now: datetime) -> bytes:
    """table は compare/table の返事。version_text は前提に書く版の表示、query は製品の絞り込みの語。"""
    openpyxl = load_openpyxl()
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill

    fill = PatternFill("solid", start_color=MARK_COLOR, end_color=MARK_COLOR)
    heads = [h["id"] for h in table["heads"]]
    multi = len(table["names"]) > 1
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(["品番", "名称", "年間台数", "", *[h["label"] for h in table["heads"]]])
    for cell in ws[1]:
        cell.font = Font(bold=True)

    def ignored(head: str) -> bool:
        return head == "AMORTIZATION" and table["ignore_amortization"]

    def put(column: int, value: str | None, *, percent: bool = False, known: bool = True, note: str | None = None,
            head: str | None = None) -> None:
        cell = ws.cell(row=ws.max_row, column=column)
        if head is not None and ignored(head):
            cell.value = DASH
            return
        if value is None:
            if percent and known:
                cell.value = DASH
                return
            cell.fill = fill
            cell.comment = Comment(note or MISSING_NOTE, "bom-cost")
            return
        number = Decimal(value)
        if percent:
            cell.value = number / 100
            cell.number_format = PERCENT_FORMAT
        else:
            cell.value = number
            cell.number_format = MONEY_FORMAT if number == number.to_integral_value() else MONEY_FRACTION_FORMAT
        if note:
            cell.comment = Comment(note, "bom-cost")

    def label(text: str, name: str | None) -> str:
        return f"{text}（{name}）" if multi and name else text

    def line(first: list[Any], cells: list[dict[str, Any]]) -> None:
        ws.append(first)
        for i, (head, spec) in enumerate(zip(heads, cells)):
            put(5 + i, head=head, **spec)

    for p in table["products"]:
        out = {h: any(p["current"][h] is None or t["after"][h] is None for t in p["trials"]) for h in heads}
        line([p["product"], p["product_name"], None, "現行"],
             [{"value": p["current"][h], "note": EXCLUDED_NOTE if out[h] else None} for h in heads])
        put(3, p["volume"], note="未登録")
        for t in p["trials"]:
            ex = {h: EXCLUDED_NOTE if p["current"][h] is None or t["after"][h] is None else None for h in heads}
            line(["", "", None, label("変更後", t["name"])], [{"value": t["after"][h], "note": ex[h]} for h in heads])
            line(["", "", None, label("差", t["name"])], [{"value": t["difference"][h], "note": ex[h]} for h in heads])
            line(["", "", None, label("差%", t["name"])],
                 [{"value": t["percent"][h], "percent": True, "known": t["difference"][h] is not None,
                   "note": ex[h]} for h in heads])
            line(["", "", None, label("年間効果額", t["name"])],
                 [{"value": t["effect"][h], "note": None if t["effect"][h] is not None else MISSING_NOTE}
                  for h in heads])

    for n in table["names"]:
        s, e = table["totals"][n], table["effect_totals"][n]
        title = label("計", n)

        def note(excluded: int) -> str | None:
            return f"未登録 {excluded} 製品を除く" if excluded else None

        line([title, "", None, "現行"], [{"value": s[h]["current"], "note": note(s[h]["excluded"])} for h in heads])
        line(["", "", None, "変更後"], [{"value": s[h]["after"], "note": note(s[h]["excluded"])} for h in heads])
        line(["", "", None, "差"], [{"value": s[h]["difference"], "note": note(s[h]["excluded"])} for h in heads])
        line(["", "", None, "差%"], [{"value": s[h]["percent"], "percent": True, "known": s[h]["difference"] is not None,
                                      "note": note(s[h]["excluded"])} for h in heads])
        line(["", "", None, "年間効果額"],
             [{"value": e[h]["effect"], "note": note(e[h]["excluded"])} for h in heads])
        for cell in ws[ws.max_row - 4][0:1]:
            cell.font = Font(bold=True)
    ws.freeze_panes = "E2"

    premises = wb.create_sheet(PREMISES)
    shown = len(table["products"]) + table["without_trial"]
    for key, value in (
        ("版", version_text),
        ("検討案", "・".join(table["names"])),
        ("製品の絞り込み", query or "なし"),
        ("製品", f"{shown} / {table['count']}" + ("（上限で打ち切り）" if table["truncated"] else "")),
        ("試算なしの製品", table["without_trial"]),
        ("償却費", "無視（—）" if table["ignore_amortization"] else "含める"),
        ("年間効果額", "差 × 年間台数"),
        ("未登録", "空欄（塗りとコメント）・計から除く"),
        ("値の扱い", "参考値"),
        ("書き出した日時", now.strftime("%Y-%m-%d %H:%M")),
    ):
        premises.append([key, value])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
